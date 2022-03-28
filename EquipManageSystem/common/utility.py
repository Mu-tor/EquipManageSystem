import os
import random
import string
import xlwt

from PIL import Image, ImageFont, ImageDraw, ImageFilter
from docx import Document
from app import app


def get_report(textlist):  # 根据模板生成借用声明
    doc = Document(os.path.join(app.static_folder, "document", "器材借用声明模板.docx"))  # 读取模板文档
    count = 0
    for p in doc.paragraphs:  # 遍历文档内容
        if 'XXXX' in p.text:  # 发现XXXX模板替换内容
            inline = p.runs
            for i in range(len(inline)):
                if 'XXXX' in inline[i].text:
                    text = inline[i].text.replace('XXXX', textlist[count])
                    inline[i].text = text  # 替换
                    count += 1
    doc.save(os.path.join(app.static_folder, "document", f"{textlist[6]}_器材借用声明.docx"))
    return doc


def get_table(datalist):  # 生成记录报表
    length = len(datalist)
    book = xlwt.Workbook()
    sheet = book.add_sheet('借用记录', cell_overwrite_ok=True)
    column = ["借用日期", "	器材名称", "	数量	", "借用人", "	归还日期	", "备注"]
    for i in range(0, 6):  # 写入列名
        sheet.write(0, i, column[i])
    for i in range(0, length):  # 获取每条数据
        data = datalist[i]
        for j in range(0, 6):
            sheet.write(i + 1, j, data[j])  # 插入每个单元格
    book.save(os.path.join(app.static_folder, "document", "器材借用记录报表.xls"))
    return book


def rndColor():
    '''随机颜色'''
    return (random.randint(32, 127), random.randint(32, 127), random.randint(32, 127))


def gene_text():
    '''生成4位验证码'''
    return ''.join(random.sample(string.ascii_letters + string.digits, 4))


def draw_lines(draw, num, width, height):
    '''划线'''
    for num in range(num):
        x1 = random.randint(0, width / 2)
        y1 = random.randint(0, height / 2)
        x2 = random.randint(0, width)
        y2 = random.randint(height / 2, height)
        draw.line(((x1, y1), (x2, y2)), fill='black', width=1)


def get_verify_code():
    '''生成验证码图形'''
    code = gene_text()
    # 图片大小120×50
    width, height = 120, 50
    # 新图片对象
    im = Image.new('RGB', (width, height), 'white')
    # 字体
    font = ImageFont.truetype('app/static/arial.ttf', 40)
    # draw对象
    draw = ImageDraw.Draw(im)
    # 绘制字符串
    for item in range(4):
        draw.text((5 + random.randint(-3, 3) + 23 * item, 5 + random.randint(-3, 3)),
                  text=code[item], fill=rndColor(), font=font)
    # 划线
    draw_lines(draw, 2, width, height)
    # 高斯模糊
    im = im.filter(ImageFilter.GaussianBlur(radius=1.5))
    return im, code
